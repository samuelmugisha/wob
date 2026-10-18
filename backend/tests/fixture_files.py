"""Builders for realistic test documents."""

import io

from PIL import Image, ImageDraw, ImageFont

LETTER = (
    "Jane Doe has served on the board of Example Holdings since 2019. As chair of the audit "
    "committee she led the remediation of internal control weaknesses identified in 2020."
)


def pdf_with_structure(encrypt: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    table = Table(
        [
            ["Year", "Organisation", "Role"],
            ["2019", "Example Holdings", "Director"],
            ["2021", "Example Holdings", "Audit Chair"],
        ]
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    story = [
        Paragraph("Nomination statement", styles["Heading1"]),
        Paragraph(LETTER, styles["Normal"]),
        Paragraph("Board appointments", styles["Heading2"]),
        table,
        PageBreak(),
        Paragraph("Community leadership", styles["Heading2"]),
        Paragraph("She mentors women directors through a regional programme.", styles["Normal"]),
    ]
    SimpleDocTemplate(buffer, pagesize=A4, encrypt=encrypt).build(story)
    return buffer.getvalue()


def text_image(text: str, fmt: str = "PNG") -> bytes:
    font = ImageFont.load_default(size=36)
    image = Image.new("RGB", (1800, 400), "white")
    ImageDraw.Draw(image).multiline_text((40, 40), text, fill="black", font=font, spacing=16)
    buffer = io.BytesIO()
    image.save(buffer, fmt)
    return buffer.getvalue()


def scanned_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    image = ImageReader(io.BytesIO(text_image(text)))
    pdf.drawImage(image, 20, 500, width=560, height=124)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def docx_document() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Supporting letter", level=1)
    document.add_paragraph(LETTER)
    document.add_heading("Board roles", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Organisation", "Role"
    table.cell(1, 0).text, table.cell(1, 1).text = "Example Holdings", "Audit Chair"
    document.add_paragraph("Signed by the Company Secretary.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_workbook() -> bytes:
    import datetime

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(["Meeting date", "Attended", None, "Notes"])
    sheet.append([datetime.datetime(2025, 3, 14), "Yes", None, "Chaired"])
    sheet.append([datetime.datetime(2025, 6, 20), "Yes", None, 4.0])
    workbook.create_sheet("Empty")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def docx_from(sections: dict[str, list[str]]) -> bytes:
    """A Word document with one level-1 heading per section and its paragraphs."""
    import docx

    document = docx.Document()
    for heading, paragraphs in sections.items():
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
