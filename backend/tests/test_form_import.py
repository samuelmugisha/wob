"""Importing nomination-form responses (a Google Forms export) as finalists' evidence."""

import io
import uuid
from datetime import datetime
from pathlib import Path

import docx
import openpyxl
import pytest
from sqlalchemy import Engine, text

from app import cli
from app.config import get_settings
from app.main import app
from app.services.form_import import match_finalist, name_key
from tests.conftest import TEST_USER_HEADER
from tests.test_knowledge_base import make_candidate

HEADERS = [
    "Timestamp",
    "Email Address",
    "Full Name",
    "Phone number (Add Country Code)",
    "Current organization",
    "Current position / Title",
    "Upload your CV or a short professional profile",
    "Referee 1 - Full name",
    "Referee 1 - Email Address",
    "Referee 1 - Mobile Number",
    "Referee 1 - Organisation and Position",
    "Declaration",
    "Date of birth",
    "Describe your exceptional leadership within your function and organization.",
]


def row(name: str, answer: str) -> list:
    return [
        datetime(2026, 9, 1, 10, 30),
        "nominee@example.org",
        name,
        "+256 700 123456",
        "Example Holdings",
        "Chief Finance Officer",
        "https://drive.google.com/open?id=abc",
        "Grace Referee",
        "grace@example.org",
        "+256 772 000111",
        "Board Chair, Example Holdings",
        "I confirm that the information is true.",
        "1990-01-01",
        answer,
    ]


def workbook(sheets: dict[str, list[list]]) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        ws = book.create_sheet(title)
        ws.append(HEADERS)
        for r in rows:
            ws.append(r)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def setup(coordinator, engine: Engine, tmp_path: Path, monkeypatch) -> dict:
    tag = uuid.uuid4().hex[:8]
    sheet = f"Chair {tag} BoardDirector Awar"[:31]  # Sheets cut titles to 31 characters
    category = coordinator.post(
        "/api/v1/categories",
        json={"code": f"chair-{tag}", "name": f"Chair {tag} BoardDirector Award"},
    ).json()
    jane = make_candidate(coordinator, category["id"], "Jane Doe")
    path = tmp_path / "evidence.xlsx"
    path.write_bytes(
        workbook(
            {
                sheet: [
                    row("Dr. Jane  DOE", "I led the finance transformation, cutting costs 20%."),
                    row("Someone Else", "Not a finalist."),
                ],
                "ESG": [],
            }
        )
    )
    monkeypatch.setattr("app.cli.get_engine", lambda: engine)
    monkeypatch.setattr("app.cli.get_settings", app.dependency_overrides[get_settings])
    return {
        "category": category,
        "jane": jane,
        "path": str(path),
        "actor": coordinator.headers[TEST_USER_HEADER],
    }


def documents(coordinator, setup) -> list[dict]:
    base = f"/api/v1/categories/{setup['category']['id']}/finalists/{setup['jane']['id']}"
    return coordinator.get(f"{base}/documents").json()


def test_preview_changes_nothing(coordinator, setup, capsys) -> None:
    cli.import_form_evidence(setup["path"], setup["actor"], confirmed=False, extract=False)
    out = capsys.readouterr().out
    assert "Dr. Jane  DOE -> Jane Doe: import 1 answers" in out
    assert "Someone Else: no matching finalist, skipped" in out
    assert "Empty sheets: ESG" in out
    assert documents(coordinator, setup) == []


def test_import_attaches_answers_without_contact_details(
    coordinator, setup, engine, capsys
) -> None:
    cli.import_form_evidence(setup["path"], setup["actor"], confirmed=True, extract=True)
    assert "Imported 1 document(s); claim extraction queued for 1 nominee(s)" in (
        capsys.readouterr().out
    )
    [doc] = documents(coordinator, setup)
    assert doc["filename"] == "Nomination form - Dr. Jane  DOE.docx"
    assert doc["source_type"] == "supporting_document"

    content = coordinator.get(f"/api/v1/documents/{doc['id']}/content").content
    body = "\n".join(p.text for p in docx.Document(io.BytesIO(content)).paragraphs)
    assert "cutting costs 20%" in body and "Chief Finance Officer" in body
    assert "Grace Referee, Board Chair, Example Holdings" in body
    assert "https://drive.google.com/open?id=abc" in body  # listed, not fetched
    for private in ("@example.org", "+256", "1990-01-01", "I confirm that"):
        assert private not in body

    # Running it again doesn't add a second copy.
    cli.import_form_evidence(setup["path"], setup["actor"], confirmed=True, extract=False)
    assert "Imported 0 document(s)" in capsys.readouterr().out
    assert len(documents(coordinator, setup)) == 1
    with engine.connect() as conn:
        assert conn.execute(
            text("SELECT count(*) FROM audit_events WHERE action = 'nomination_form.imported'")
        ).scalar_one()


def test_only_coordinators_can_import(client_for, setup) -> None:
    judge = client_for("judge").headers[TEST_USER_HEADER]
    with pytest.raises(SystemExit, match="not allowed to manage finalists and evidence"):
        cli.import_form_evidence(setup["path"], judge, confirmed=True, extract=False)


def test_names_match_despite_titles_case_and_accents() -> None:
    assert name_key("Dr. Jane  DOE") == name_key("jane doe") == "jane doe"
    assert name_key("Prof Ámara Okello") == "amara okello"

    class F:
        def __init__(self, name: str) -> None:
            self.full_name = name

    finalists = [(None, F("Jane Doe")), (None, F("Janet Dee"))]
    assert match_finalist("Mrs Jane Doe", finalists)[1].full_name == "Jane Doe"
    assert match_finalist("Mary Major", finalists) is None


def test_reordered_names_and_word_subset_sheets_match() -> None:
    from app.services.form_import import match_category

    class F:
        def __init__(self, name: str) -> None:
            self.full_name = name

    finalists = [(None, F("Nabirye Sarah")), (None, F("Hajjat Aphwa Sebyala Kawaase"))]
    assert match_finalist("SARAH NABIRYE", finalists)[1].full_name == "Nabirye Sarah"
    assert match_finalist("HAJJAT  APHWA KAWAASE  SEBYALA", finalists) is finalists[1]

    class C:
        def __init__(self, name: str) -> None:
            self.id, self.name = uuid.uuid4(), name

    mentorship = C("Mentorship Excellence Award – Individual")
    categories = [mentorship, C("Trailblazer Award"), C("Woman CEO of the Year")]
    assert match_category("Mentorship Award", categories) is mentorship
    assert match_category("Female Board ChairpersonDirecto", categories) is None


def test_a_person_can_map_sheets_and_names(coordinator, setup, tmp_path, capsys) -> None:
    path = tmp_path / "other.xlsx"
    path.write_bytes(workbook({"Chair Of The Year": [row("J. Doe", "I chaired the board.")]}))
    cli.import_form_evidence(str(path), setup["actor"], confirmed=False, extract=False)
    assert "Sheets with no matching category: Chair Of The Year" in capsys.readouterr().out

    cli.import_form_evidence(
        str(path),
        setup["actor"],
        confirmed=False,
        extract=False,
        sheets=[f"Chair Of The Year={setup['category']['code']}"],
        names=["J. Doe=Jane Doe"],
    )
    assert "J. Doe -> Jane Doe: import 1 answers" in capsys.readouterr().out

    with pytest.raises(SystemExit, match="no visible category has the code"):
        cli.import_form_evidence(
            str(path), setup["actor"], False, False, sheets=["Chair Of The Year=nope"]
        )
    with pytest.raises(SystemExit, match="takes LEFT=RIGHT"):
        cli.import_form_evidence(str(path), setup["actor"], False, False, names=["J. Doe"])
