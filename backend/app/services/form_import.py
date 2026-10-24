"""Import nomination-form responses (a Google Forms export, one sheet per award category) as
evidence documents for the finalists already in each category.

Each matched response becomes a Word document of the nominee's organisation and title, every
written answer under its question, and their referees' names and organisations. Contact
details (emails, phone numbers) and date of birth are left out: they are not evidence, and the
documents go to the AI and to research. Uploaded-file answers are Google Drive links the server
can't open, so they are listed for a coordinator to fetch. Responses that match no finalist in
their category are skipped.
"""

import difflib
import hashlib
import io
import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

import docx
import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.models import AwardCategory, Document, Finalist, FinalistCategory, SourceType, User
from app.services import audit, consent
from app.services.documents import Upload, store_uploads
from app.services.storage import LocalStorage

# Columns that are personal contact data or not evidence.
EXCLUDED = re.compile(r"e-?mail|phone|mobile|date of birth|^declaration$|^timestamp$", re.I)
REFEREE = re.compile(r"^referee\s*(\d+)\s*-\s*(full name|organisation and position)", re.I)
PROFILE = {"current organization": "Organisation", "current position / title": "Position"}
TITLES = re.compile(r"^(dr|mrs|mr|ms|miss|prof|eng|hon|rev)\.?\s+", re.I)


def _key(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def name_key(name: str) -> str:
    name = _key(name)
    while TITLES.match(name):
        name = TITLES.sub("", name)
    return name


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d %B %Y")
    return str(value).strip()


@dataclass
class Response:
    sheet: str
    row: int
    full_name: str
    submitted: str
    organisation: str
    position: str
    answers: list[tuple[str, str]] = field(default_factory=list)
    referees: list[str] = field(default_factory=list)
    links: list[tuple[str, str]] = field(default_factory=list)


def read_workbook(data: bytes) -> dict[str, list[Response]]:
    """Responses per sheet, with contact details already dropped."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: dict[str, list[Response]] = {}
    for ws in workbook.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_text(h) for h in rows[0]]
        responses = []
        for number, row in enumerate(rows[1:], start=2):
            cells = dict(zip(headers, row, strict=False))
            name = _text(cells.get("Full Name"))
            if not name:
                continue
            response = Response(
                sheet=ws.title,
                row=number,
                full_name=name,
                submitted=_text(cells.get("Timestamp")),
                organisation=_text(cells.get("Current organization")),
                position=_text(cells.get("Current position / Title")),
            )
            referees: dict[str, dict[str, str]] = {}
            for header, value in zip(headers, row, strict=False):
                answer = _text(value)
                if not header or not answer or header == "Full Name":
                    continue
                if header.lower() in PROFILE or EXCLUDED.search(header):
                    continue
                if match := REFEREE.match(header):
                    referees.setdefault(match.group(1), {})[match.group(2).lower()] = answer
                elif answer.startswith(("http://", "https://")):
                    response.links.append((header, answer))
                else:
                    response.answers.append((header, answer))
            for _, ref in sorted(referees.items()):
                if ref.get("full name"):
                    detail = ref.get("organisation and position")
                    response.referees.append(ref["full name"] + (f", {detail}" if detail else ""))
            responses.append(response)
        sheets[ws.title] = responses
    return sheets


def match_category(sheet: str, categories: list[AwardCategory]) -> AwardCategory | None:
    """Sheet titles are cut to 31 characters, so a category whose name starts with the title
    matches; otherwise the one category containing all its words, or the one closely similar
    name."""
    key = _key(sheet).replace(" ", "")
    names = {c.id: _key(c.name).replace(" ", "") for c in categories}
    starts = [c for c in categories if names[c.id].startswith(key) or key.startswith(names[c.id])]
    if len(starts) == 1:
        return starts[0]
    # Every word of the sheet title in one category's name ("Mentorship Award" ->
    # "Mentorship Excellence Award - Individual").
    words = set(_key(sheet).split())
    within = [c for c in categories if words <= set(_key(c.name).split())]
    if len(within) == 1:
        return within[0]
    close = [
        c for c in categories if difflib.SequenceMatcher(None, key, names[c.id]).ratio() >= 0.85
    ]
    return close[0] if len(close) == 1 else None


def match_finalist(
    name: str, finalists: list[tuple[FinalistCategory, Finalist]]
) -> tuple[FinalistCategory, Finalist] | None:
    """The finalist with the same name (ignoring case, accents, punctuation and titles), else
    the same words in another order ("Sarah Nabirye" / "Nabirye Sarah"), else the one name that
    is almost identical."""
    key = name_key(name)
    exact = [f for f in finalists if name_key(f[1].full_name) == key]
    if len(exact) == 1:
        return exact[0]
    words = sorted(key.split())
    reordered = [f for f in finalists if sorted(name_key(f[1].full_name).split()) == words]
    if len(reordered) == 1:
        return reordered[0]
    close = [
        f
        for f in finalists
        if difflib.SequenceMatcher(None, key, name_key(f[1].full_name)).ratio() >= 0.9
    ]
    return close[0] if len(close) == 1 else None


def document_name(response: Response) -> str:
    return f"Nomination form - {response.full_name}.docx"


def build_document(response: Response, category: str) -> bytes:
    document = docx.Document()
    document.core_properties.title = f"Nomination form: {response.full_name}"
    document.add_heading(f"Nomination form responses: {response.full_name}", level=0)
    document.add_paragraph(
        f"Award category: {category}."
        + (f" Submitted {response.submitted}." if response.submitted else "")
    )
    profile = [
        f"{label}: {value}"
        for label, value in (
            ("Organisation", response.organisation),
            ("Position", response.position),
        )
        if value
    ]
    if profile:
        document.add_heading("Current role", level=1)
        for line in profile:
            document.add_paragraph(line)
    for question, answer in response.answers:
        document.add_heading(question, level=1)
        for paragraph in re.split(r"\n\s*\n", answer):
            document.add_paragraph(paragraph.strip())
    if response.referees:
        document.add_heading("Referees", level=1)
        for referee in response.referees:
            document.add_paragraph(referee)
    if response.links:
        document.add_heading("Files uploaded with the form (not imported)", level=1)
        for question, link in response.links:
            document.add_paragraph(f"{question}: {link}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@dataclass
class Planned:
    response: Response
    category: AwardCategory
    finalist_category: FinalistCategory
    finalist: Finalist
    already_imported: bool
    blocked: str | None = None  # why it can't be imported (consent)


@dataclass
class Plan:
    imports: list[Planned] = field(default_factory=list)
    unmatched: list[Response] = field(default_factory=list)
    sheets_without_category: list[str] = field(default_factory=list)
    empty_sheets: list[str] = field(default_factory=list)


def plan_import(
    session: Session,
    sheets: dict[str, list[Response]],
    settings: Settings,
    sheet_categories: dict[str, str] | None = None,
    names: dict[str, str] | None = None,
) -> Plan:
    """sheet_categories maps a sheet title to a category code, and names a form's full name to
    a finalist's, for the cases a person decides; both override automatic matching."""
    plan = Plan()
    categories = session.scalars(
        select(AwardCategory).where(AwardCategory.hidden_at.is_(None))
    ).all()
    by_code = {c.code: c for c in categories}
    sheet_categories = sheet_categories or {}
    names = {name_key(k): v for k, v in (names or {}).items()}
    for code in sheet_categories.values():
        if code not in by_code:
            raise LookupError(f"no visible category has the code {code!r}")
    for sheet, responses in sheets.items():
        if not responses:
            plan.empty_sheets.append(sheet)
            continue
        if sheet in sheet_categories:
            category = by_code[sheet_categories[sheet]]
        else:
            category = match_category(sheet, list(categories))
        if category is None:
            plan.sheets_without_category.append(sheet)
            continue
        finalists = session.execute(
            select(FinalistCategory, Finalist)
            .join(Finalist, Finalist.id == FinalistCategory.finalist_id)
            .where(FinalistCategory.category_id == category.id)
        ).all()
        for response in responses:
            found = match_finalist(
                names.get(name_key(response.full_name), response.full_name),
                [tuple(f) for f in finalists],
            )
            if found is None:
                plan.unmatched.append(response)
                continue
            fc, finalist = found
            already = session.scalar(
                select(Document.id).where(
                    Document.finalist_category_id == fc.id,
                    Document.filename == document_name(response),
                    Document.withdrawn_at.is_(None),
                )
            )
            blocked = None
            if consent.withdrawn(session, fc.id, consent.Purpose.ASSESSMENT):
                blocked = "withdrew consent to assessment (part A)"
            elif settings.upload_requires_consent and not consent.has(
                session, fc.id, consent.Purpose.ASSESSMENT
            ):
                blocked = "no consent to assessment (part A) recorded"
            plan.imports.append(
                Planned(response, category, fc, finalist, already is not None, blocked)
            )
    return plan


def apply_import(
    session: Session, storage: LocalStorage, settings: Settings, plan: Plan, user: User
) -> list[Planned]:
    """Store a document for each matched nominee not imported before, and fill in their
    organisation and title where the system has none. Returns what was imported."""
    done = []
    for item in plan.imports:
        if item.already_imported or item.blocked:
            continue
        data = build_document(item.response, item.category.name)
        result = store_uploads(
            session,
            storage,
            settings,
            finalist_category=item.finalist_category,
            user=user,
            uploads=[Upload(document_name(item.response), data)],
            source_type=SourceType.SUPPORTING_DOCUMENT,
            idempotency_key=hashlib.sha256(
                f"form-import:{item.finalist_category.id}:{item.response.row}".encode() + data
            ).hexdigest(),
        )
        if not result.accepted:
            continue
        if not item.finalist.organisation and item.response.organisation:
            item.finalist.organisation = item.response.organisation[:200]
        if not item.finalist.position_title and item.response.position:
            item.finalist.position_title = item.response.position[:200]
        audit.record(
            session,
            user,
            "nomination_form.imported",
            "finalist_category",
            item.finalist_category.id,
            sheet=item.response.sheet,
            row=item.response.row,
            answers=len(item.response.answers),
            document_id=str(result.accepted[0].id),
        )
        done.append(item)
    return done


def finalist_category_ids(items: list[Planned]) -> list[uuid.UUID]:
    return [i.finalist_category.id for i in items]
